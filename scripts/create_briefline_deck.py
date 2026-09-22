from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = "briefline-short-pitch.pptx"
W, H = 13.333, 7.5
BG = RGBColor(7, 20, 35)
PANEL = RGBColor(19, 40, 68)
LINE = RGBColor(41, 69, 104)
WHITE = RGBColor(238, 245, 255)
MUTED = RGBColor(169, 189, 215)
BLUE = RGBColor(62, 139, 255)
CYAN = RGBColor(99, 230, 190)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.12
    return shape

def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, font="Aptos",
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box

def base():
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, BG)
    # Decorative accents
    rect(s, 11.55, -0.55, 2.4, 2.4, RGBColor(18, 63, 130), line=RGBColor(18,63,130), radius=True)
    rect(s, 11.95, -0.18, 1.4, 1.4, RGBColor(36, 105, 194), line=RGBColor(36,105,194), radius=True)
    return s

def footer(s, label, page):
    text(s, label.upper(), .72, 7.13, 6.8, .18, 8, RGBColor(105,131,163), True)
    text(s, f"{page:02d} / 07", 11.7, 7.13, .9, .18, 8, RGBColor(105,131,163), True, align=PP_ALIGN.RIGHT)

def kicker(s, value):
    text(s, value.upper(), .72, .58, 7, .22, 10, CYAN, True)

def title(s, value, y=.98, size=29, h=1.0):
    text(s, value, .72, y, 11.8, h, size, WHITE, True)

def card(s, x, y, w, h, heading, body, icon=None, fill=PANEL):
    rect(s, x, y, w, h, fill, line=LINE, radius=True)
    if icon:
        rect(s, x+.22, y+.22, .42, .42, RGBColor(29,92,184), line=RGBColor(29,92,184), radius=True)
        text(s, icon, x+.22, y+.23, .42, .34, 9, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        top = y+.83
    else:
        top = y+.24
    text(s, heading, x+.22, top, w-.44, .28, 13, WHITE, True)
    text(s, body, x+.22, top+.38, w-.44, h-(top-y)-.55, 10, MUTED)

# 1
s = base(); kicker(s, "Snapdragon AI PC Challenge")
text(s, "Meetings into", .72, 1.23, 8.8, .62, 40, WHITE, True)
text(s, "action.", .72, 1.84, 8.8, .62, 40, CYAN, True)
text(s, "Private by design.", .72, 2.48, 10, .55, 30, WHITE, True)
text(s, "Briefline is an evidence-linked meeting and document copilot built for Snapdragon-powered HP PCs.", .72, 3.32, 8.8, .75, 18, RGBColor(219,233,250))
for i, label in enumerate(["On-device AI", "Offline-first", "Accessible", "Evidence-linked"]):
    rect(s, .72+i*1.62, 4.52, 1.48 if i<2 else 1.65, .36, BG, line=LINE, radius=True)
    text(s, label, .72+i*1.62, 4.57, 1.48 if i<2 else 1.65, .22, 9, RGBColor(207,224,245), True, align=PP_ALIGN.CENTER)
footer(s, "Briefline", 1)

# 2
s = base(); kicker(s, "The problem"); title(s, "The most important part of a meeting is what happens after it.", .98, 29, 1.16)
card(s, .72, 2.65, 3.85, 2.15, "Context gets lost", "People remember fragments, not the full chain from conversation to decision.", "01")
card(s, 4.73, 2.65, 3.85, 2.15, "Cloud is not always acceptable", "Confidential client work, unpublished ideas, and weak connectivity need a local option.", "02")
card(s, 8.74, 2.65, 3.85, 2.15, "Summaries are hard to trust", "A polished paragraph is not enough if nobody can check where it came from.", "03")
footer(s, "Why now", 2)

# 3
s = base(); kicker(s, "The solution"); title(s, "Briefline creates a brief you can verify.", .98, 32, .55)
text(s, "It listens locally, structures what changed, and links every action or decision to the source that supports it.", .72, 1.72, 7.0, .7, 17, MUTED)
rect(s, .72, 2.82, 7.05, 1.05, BG, line=CYAN, radius=True)
text(s, "“Useful conclusion + the evidence behind it — without uploading the work.”", .98, 3.04, 6.55, .58, 17, WHITE, False, italic=True)
card(s, 8.35, 1.65, 4.22, 3.4, "One local workspace", "Live transcript\nDecisions and action owners\nDeadlines and open questions\nTimestamp and document evidence\nEditable export-ready brief", "↗")
footer(s, "What we build", 3)

# 4
s = base(); kicker(s, "How it works"); title(s, "From conversation to a checked next step.", .98, 31, .6)
items = [
    ("01 / CAPTURE", "Microphone or files", "Audio, PDFs, images, and notes remain inside the local workspace."),
    ("02 / UNDERSTAND", "Transcribe and read", "Local models extract speech and document context."),
    ("03 / STRUCTURE", "Find what changed", "Decisions, owners, dates, and questions become editable cards."),
    ("04 / VERIFY", "Show the evidence", "Every result opens its supporting timestamp or source region."),
]
for i, (num, head, body) in enumerate(items):
    card(s, .72+i*3.15, 2.05, 2.87, 2.7, head, body, num.split(" ")[0])
footer(s, "Product flow", 4)

# 5
s = base(); kicker(s, "Built for Snapdragon"); title(s, "Use the device’s AI engine, not a generic cloud endpoint.", .98, 27, .7)
card(s, .72, 2.05, 3.32, 2.2, "Qualcomm AI Hub", "Whisper-Medium for multilingual speech recognition.\n\nQwen3-0.6B for structured extraction.", "01")
text(s, "→", 4.2, 2.78, .5, .45, 27, CYAN, True, align=PP_ALIGN.CENTER)
card(s, 4.95, 2.05, 3.45, 2.2, "Snapdragon AI Engine", "Profile and deploy compiled assets for NPU/GPU/CPU execution on the HP PC.", "02", fill=RGBColor(22,55,101))
text(s, "→", 8.58, 2.78, .5, .45, 27, CYAN, True, align=PP_ALIGN.CENTER)
card(s, 9.3, 2.05, 3.27, 2.2, "Briefline app", "Tauri + React interface, local SQLite search, offline export.", "03")
for i, (big, small) in enumerate([("0","required cloud calls in the core flow"),("100%","of results designed to carry evidence"),("2 tiers","Snapdragon X Elite and X Plus"),("3 modes","quality, balanced, low-power")]):
    x=.72+i*3.15
    text(s, big, x, 5.05, 2.4, .45, 25, CYAN, True)
    text(s, small, x, 5.57, 2.4, .45, 9, MUTED)
footer(s, "Technical edge", 5)

# 6
s = base(); kicker(s, "Trust and access"); title(s, "Private enough for real work. Simple enough for everyone.", .98, 28, .8)
checks = ["Local-only mode enabled by default","Pause and delete controls per session","No account required for core workflow","Keyboard navigation and visible focus","Scalable typography and high contrast","Uncertain output is labeled, not hidden"]
for i, item in enumerate(checks):
    x = .72 + (i%2)*6.0
    y = 2.08 + (i//2)*.75
    rect(s, x, y, 5.55, .5, PANEL, line=LINE, radius=True)
    text(s, "✓", x+.18, y+.09, .27, .22, 13, CYAN, True)
    text(s, item, x+.57, y+.11, 4.75, .22, 11, RGBColor(217,231,247))
rect(s, .72, 4.8, 10.5, .85, BG, line=CYAN, radius=True)
text(s, "A local AI assistant should not just be faster. It should make the user more comfortable using it.", .98, 5.02, 10.0, .4, 17, WHITE, False, italic=True)
footer(s, "Deployment & accessibility", 6)

# 7
s = base(); kicker(s, "The finish line")
text(s, "Make the next step obvious.", .72, 1.18, 10.6, .62, 35, WHITE, True)
text(s, "Keep the work yours.", .72, 1.85, 10.6, .62, 35, CYAN, True)
text(s, "Briefline turns the Snapdragon-powered HP PC into a private, practical partner for the work that follows every conversation.", .72, 2.83, 9.3, .75, 18, RGBColor(219,233,250))
for i, label in enumerate(["Working prototype", "Measured on device", "Open implementation"]):
    rect(s, .72+i*2.1, 4.25, 1.9, .36, BG, line=LINE, radius=True)
    text(s, label, .72+i*2.1, 4.31, 1.9, .2, 9, RGBColor(207,224,245), True, align=PP_ALIGN.CENTER)
footer(s, "Briefline · Snapdragon AI PC Challenge", 7)

prs.save(OUT)
print(OUT)